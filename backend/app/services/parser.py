import re
from pathlib import Path

from pypdf import PdfReader
from docx import Document

ALLOWED_EXTENSIONS = {".pdf", ".docx"}


class ResumeParseError(Exception):
    pass


def extract_text(file_stream, filename: str) -> str:
    ext = Path(filename).suffix.lower()

    if ext == ".pdf":
        text = _from_pdf(file_stream)
    elif ext == ".docx":
        text = _from_docx(file_stream)
    else:
        raise ResumeParseError(f"Unsupported file type: {ext}")

    text = _clean(text)

    if len(text) < 100:
        raise ResumeParseError(
            "Could not extract enough text. The file may be scanned or image-based."
        )

    return text


def _from_pdf(stream) -> str:
    reader = PdfReader(stream)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _from_docx(stream) -> str:
    doc = Document(stream)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _clean(text: str) -> str:
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()